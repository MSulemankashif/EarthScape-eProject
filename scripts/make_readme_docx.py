"""Generate ReadMe.docx (Word deliverable) from README.md.

Run:  python scripts/make_readme_docx.py
Safe to re-run whenever README.md changes.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "README.md"
DST = ROOT / "ReadMe.docx"

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_rich(par, text: str) -> None:
    """Append text to a paragraph, rendering **bold**, `code` and [links](url)."""
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        seg = m.group(0)
        if seg.startswith("**"):
            run = par.add_run(seg[2:-2])
            run.bold = True
        elif seg.startswith("`"):
            run = par.add_run(seg[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x93, 0x33, 0x18)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", seg)
            run = par.add_run(f"{link.group(1)} ({link.group(2)})")
            run.font.color.rgb = RGBColor(0x1A, 0x5F, 0xB4)
            run.underline = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def main() -> None:
    doc = Document()
    doc.core_properties.title = "EarthScape Climate Agency - README"
    doc.core_properties.author = "EarthScape Climate Agency"

    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)

    lines = SRC.read_text(encoding="utf-8").splitlines()

    in_code = False
    code_buf: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]

        # ---- fenced code blocks -------------------------------------------------
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x3D)
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # ---- headings -----------------------------------------------------------
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = doc.add_heading(level=min(level, 4))
            add_rich(heading, line.lstrip("#").strip())

        # ---- tables (collect contiguous | rows) ---------------------------------
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci in range(width):
                        text = row[ci] if ci < len(row) else ""
                        cell_p = table.cell(ri, ci).paragraphs[0]
                        add_rich(cell_p, text)
                        if ri == 0:
                            for r in cell_p.runs:
                                r.bold = True
                doc.add_paragraph()
            continue

        # ---- lists --------------------------------------------------------------
        elif line.startswith("- ") or re.match(r"\d+\. ", line):
            style = "List Bullet" if line.startswith("- ") else "List Number"
            p = doc.add_paragraph(style=style)
            add_rich(p, re.sub(r"^([-*]|\d+\.) ", "", line))

        # ---- horizontal rules ---------------------------------------------------
        elif line.strip() == "---":
            p = doc.add_paragraph()
            rule = p.add_run("─" * 74)
            rule.font.color.rgb = RGBColor(0xBB, 0xC4, 0xD4)

        elif line.strip():
            p = doc.add_paragraph()
            add_rich(p, line)

        i += 1

    doc.save(DST)
    print(f"{DST.name} written: {len(doc.paragraphs)} paragraphs, "
          f"{len(doc.tables)} tables")


if __name__ == "__main__":
    main()
